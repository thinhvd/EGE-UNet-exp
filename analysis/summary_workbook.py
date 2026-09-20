'''
One-page summary workbook: accuracy and cost of every trained variant, side by side.

Reads only files produced by the rest of the pipeline - test_results.json per run (pooled metrics),
the stratified comparison of compare_runs.py (per-image means, overall and per lesion-size group),
and the cost table of benchmark_speed.py - so nothing is retyped. Derived cells (GFLOPs, per-image
latency, throughput) are Excel formulas over the measured cells.

Usage:
  python analysis/summary_workbook.py \
      --run learnable=results/.../egeunet_isic17_learnable_s42 [--run ...] \
      --strata results/.../compare_fusion10_dsc/compare_runs_dsc_strata.csv \
      --cost results/.../cost_table_all10.csv \
      --out results/.../tong_hop_ket_qua.xlsx
The --run labels must match the group names in the strata csv and the model names in the cost csv.
'''
import os
import sys
import csv
import json
import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=13, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color='52514E')
FILL_HEAD = PatternFill('solid', fgColor='D6E4F0')
FILL_REF = PatternFill('solid', fgColor='EEF3F8')
THIN = Side(style='thin', color='C3C2B7')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical='top')

COLS = [
    ('Model', 22, None),
    ('Tham số', 10, '#,##0'),
    ('GMACs', 9, '0.0000'),
    ('GFLOPs', 9, '0.0000'),
    ('mIoU toàn tập test (%)', 12, '0.00'),
    ('DSC toàn tập test (%)', 12, '0.00'),
    ('DSC trung bình theo ảnh (%)', 13, '0.00'),
    ('DSC nhóm lesion NHỎ (%)', 13, '0.00'),
    ('DSC nhóm lesion TO (%)', 13, '0.00'),
    ('GPU: ms cho 1 ảnh (lô 1)', 12, '0.00'),
    ('GPU: ms cho cả lô 8', 12, '0.00'),
    ('GPU: ms/ảnh khi chạy lô 8', 13, '0.00'),
    ('GPU: ảnh/giây (lô 8)', 12, '0'),
    ('CPU: ms cho 1 ảnh (lô 1)', 12, '0.00'),
    ('CPU: ms cho cả lô 8', 12, '0.00'),
    ('CPU: ms/ảnh khi chạy lô 8', 13, '0.00'),
    ('CPU: ảnh/giây (lô 8)', 12, '0'),
]
# columns where a higher value is better / a lower value is better, for the bold-best marking
BEST_MAX = {5, 6, 7, 8, 9, 13, 17}
BEST_MIN = {2, 3, 4, 10, 11, 12, 14, 15, 16}


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--run', action='append', required=True, metavar='LABEL=DIR')
    p.add_argument('--strata', required=True, help='compare_runs *_strata.csv (metric dsc)')
    p.add_argument('--cost', required=True, help='cost table csv from the benchmark')
    p.add_argument('--baseline', default='learnable', help='row highlighted as the reference model')
    p.add_argument('--exclude-from-best', default='none',
                   help='comma list of labels that are ablations, not candidates (not marked best)')
    p.add_argument('--out', required=True)
    return p.parse_args()


def put(ws, r, c, v, font=F_BODY, fill=None, fmt=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    cell.border = BOX
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    return cell


def main():
    a = parse_args()
    runs = [s.split('=', 1) for s in a.run]
    strata = {(r['stratum'], r['group']): r for r in csv.DictReader(open(a.strata))
              if r['stratification'] == 'area_frac'}
    cost = {r['model']: r for r in csv.DictReader(open(a.cost))}
    small = next(k[0] for k in strata if k[0].startswith('area_T1'))
    large = next(k[0] for k in strata if k[0].startswith('area_T3'))

    missing = [lab for lab, _ in runs if lab not in cost or (small, lab) not in strata]
    if missing:
        raise SystemExit(f'no cost row / strata row for: {missing}')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tổng hợp'
    put(ws, 1, 1, 'EGE-UNet — tổng hợp kết quả và chi phí của các biến thể đã train', F_TITLE).border = None
    put(ws, 2, 1, 'Cột nền xanh nhạt: model gốc của bài báo. Số in đậm: tốt nhất trong các biến thể '
                  '(không tính model gốc và bản bỏ GHPA). Xem sheet "Ghi chú" để biết nguồn và cách đo.',
        F_NOTE).border = None

    head = 4
    for j, (title, width, _) in enumerate(COLS, 1):
        put(ws, head, j, title, F_BOLD, FILL_HEAD, align=CENTER)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[head].height = 56

    values = {}
    for i, (lab, d) in enumerate(runs):
        row = head + 1 + i
        t = json.load(open(os.path.join(d, 'test_results.json')))
        c = cost[lab]
        s_all = strata[('all', lab)]['mean']
        fill = FILL_REF if lab == a.baseline else None
        put(ws, row, 1, lab, F_BOLD if lab == a.baseline else F_BODY, fill)
        put(ws, row, 2, int(c['params']), fill=fill, fmt=COLS[1][2])
        put(ws, row, 3, float(c['gmacs']), fill=fill, fmt=COLS[2][2])
        put(ws, row, 4, f'=2*C{row}', fill=fill, fmt=COLS[3][2])
        put(ws, row, 5, 100 * t['miou'], fill=fill, fmt=COLS[4][2])
        put(ws, row, 6, 100 * t['f1_or_dsc'], fill=fill, fmt=COLS[5][2])
        put(ws, row, 7, 100 * float(s_all), fill=fill, fmt=COLS[6][2])
        put(ws, row, 8, 100 * float(strata[(small, lab)]['mean']), fill=fill, fmt=COLS[7][2])
        put(ws, row, 9, 100 * float(strata[(large, lab)]['mean']), fill=fill, fmt=COLS[8][2])
        put(ws, row, 10, float(c['gpu_bs1_ms']), fill=fill, fmt=COLS[9][2])
        put(ws, row, 11, float(c['gpu_bs8_ms']), fill=fill, fmt=COLS[10][2])
        put(ws, row, 12, f'=K{row}/8', fill=fill, fmt=COLS[11][2])
        put(ws, row, 13, f'=1000/L{row}', fill=fill, fmt=COLS[12][2])
        put(ws, row, 14, float(c['cpu_bs1_ms']), fill=fill, fmt=COLS[13][2])
        put(ws, row, 15, float(c['cpu_bs8_ms']), fill=fill, fmt=COLS[14][2])
        put(ws, row, 16, f'=O{row}/8', fill=fill, fmt=COLS[15][2])
        put(ws, row, 17, f'=1000/P{row}', fill=fill, fmt=COLS[16][2])
        values[lab] = {
            2: int(c['params']), 3: float(c['gmacs']), 4: 2 * float(c['gmacs']),
            5: 100 * t['miou'], 6: 100 * t['f1_or_dsc'], 7: 100 * float(s_all),
            8: 100 * float(strata[(small, lab)]['mean']), 9: 100 * float(strata[(large, lab)]['mean']),
            10: float(c['gpu_bs1_ms']), 11: float(c['gpu_bs8_ms']), 12: float(c['gpu_bs8_ms']) / 8,
            13: 8000 / float(c['gpu_bs8_ms']), 14: float(c['cpu_bs1_ms']), 15: float(c['cpu_bs8_ms']),
            16: float(c['cpu_bs8_ms']) / 8, 17: 8000 / float(c['cpu_bs8_ms']),
        }

    skip = {a.baseline} | {x for x in a.exclude_from_best.split(',') if x}
    cands = [lab for lab, _ in runs if lab not in skip]
    for col in BEST_MAX | BEST_MIN:
        pick = (max if col in BEST_MAX else min)(cands, key=lambda l: values[l][col])
        r = head + 1 + [lab for lab, _ in runs].index(pick)
        ws.cell(row=r, column=col).font = F_BOLD

    ws.freeze_panes = f'B{head + 1}'
    last = head + len(runs)
    ws.auto_filter.ref = f'A{head}:{get_column_letter(len(COLS))}{last}'

    notes = wb.create_sheet('Ghi chú')
    notes.column_dimensions['A'].width = 34
    notes.column_dimensions['B'].width = 110
    put(notes, 1, 1, 'Nguồn số liệu và cách đo', F_TITLE).border = None
    items = [
        ('mIoU / DSC toàn tập test',
         'Gộp toàn bộ pixel của 650 ảnh val rồi tính một con số (đúng cách engine.py của tác giả tính). '
         'Lấy từ test_results.json của từng run — chính là con số dùng trong báo cáo.'),
        ('DSC trung bình theo ảnh',
         'Tính DSC cho từng ảnh rồi lấy trung bình 650 ảnh, nên mọi ảnh có trọng số như nhau. '
         'Khác với DSC toàn tập ở trên, vốn bị ảnh tổn thương to chi phối.'),
        ('DSC nhóm lesion NHỎ / TO',
         'Chia 650 ảnh thành 3 nhóm bằng nhau theo diện tích tổn thương; NHỎ là 217 ảnh có tổn thương '
         'dưới 5,45% diện tích ảnh, TO là 217 ảnh trên 16,53%. Giá trị là trung bình DSC theo ảnh trong nhóm.'),
        ('Tham số / GMACs',
         'Đại lượng tất định. Đã đo độc lập trên CPU local và trên GPU (3 vòng) và trên một máy GPU khác — '
         'cả ba cho kết quả trùng khớp tuyệt đối. GMACs đếm mọi conv/linear cộng phần einsum của attention; '
         'không đếm chuẩn hoá, GELU, nội suy, softmax, nên là cận dưới — nhưng thiếu giống nhau ở mọi biến thể.'),
        ('GMACs so với "GFLOPs" của bài báo',
         'Bài báo EGE-UNet ghi 0,072 GFLOPs. Con số đó chính là cột GMACs ở đây (công cụ đo phổ biến gọi MACs '
         'là FLOPs); model gốc đo được 0,0721 GMACs. Cột GFLOPs = 2 × GMACs, tức số phép tính thật.'),
        ('Tốc độ',
         'Ảnh 256×256, chế độ suy luận (không tính gradient). Mỗi biến thể đo 3 vòng, thứ tự các biến thể '
         'xáo lại mỗi vòng; giá trị ghi ở đây là trung vị của 3 vòng. GPU: RTX 4090 thuê, 200 lần lặp/vòng '
         'sau 50 lần làm nóng. CPU: máy local, 4 luồng, 50 lần lặp/vòng.'),
        ('Đọc số tốc độ thế nào',
         'Chênh lệch dưới 1 ms trên GPU KHÔNG có ý nghĩa: cùng một biến thể dao động tới ±1 ms giữa các vòng. '
         'Trên GPU, thứ tự tốc độ không bám theo GMACs vì model quá nhỏ nên thời gian chủ yếu là chi phí khởi '
         'chạy kernel; trên CPU thì bám sát hơn. Latency của hai máy khác nhau không so trực tiếp được.'),
        ('Điều kiện train',
         'ISIC2017, 1500 ảnh train / 650 ảnh val, 256×256, 300 epoch, batch 8, AdamW lr 1e-3, seed 42, '
         'mỗi cấu hình 1 lần train. Tập val đồng thời dùng để chọn model và để chấm điểm (theo code gốc).'),
        ('Lưu ý quan trọng khi so sánh',
         'Mỗi lần train là một lần rút ngẫu nhiên: góc xoay augmentation được rút lúc import, trước khi đặt seed, '
         'và bước backward trên GPU không tất định. Cùng một cấu hình đã đo được lệch khoảng 1 điểm mIoU giữa '
         'hai lần train. Vì vậy chênh lệch dưới ~1 điểm mIoU giữa hai model, với 1 seed, chưa đủ để kết luận.'),
    ]
    for i, (k, v) in enumerate(items, 3):
        put(notes, i, 1, k, F_BOLD, align=WRAP)
        put(notes, i, 2, v, align=WRAP)
        notes.row_dimensions[i].height = 42

    wb.save(a.out)
    print(f'wrote {a.out}')


if __name__ == '__main__':
    sys.exit(main())
