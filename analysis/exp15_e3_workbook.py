'''EXP-15: results workbook for the E3 report (LSA-CSF route + probe vs the original model and CSF), same layout as
results/EGE-UNet-results-exp10/exp10_tong_hop.xlsx. Every number is recomputed from the per-image test csv files.
  python analysis/exp15_e3_workbook.py   -> results/EGE-UNet-results-exp15/exp15_E3_tong_hop.xlsx
'''
import csv, os, numpy as np, openpyxl
from scipy import stats
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
R = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'results')
OUT = f'{R}/EGE-UNet-results-exp15/exp15_E3_tong_hop.xlsx'
ref = openpyxl.load_workbook(f'{R}/EGE-UNet-results-exp10/exp10_tong_hop.xlsx')['isic17']
FMT_T, FMT_POOL = ref['D19'].number_format, '0.00'
FMT_D, FMT_P = ref['C19'].number_format, ref['E19'].number_format
CSV = 'analysis/per_image_metrics_test.csv'
ARMS = [('base', 'Model gốc (EGE-UNet)', 'learnable', 53374),
        ('csf', 'Model gốc + CSF ở dec3–5 (8 kênh)', 'learnable_fuse-sum-shallow3-d8', 54965),
        ('lsa', 'Model gốc + LSA-CSF ở dec3–5 (router chọn lọc)', 'learnable_fuse-sum-shallow3-d8_lsa-route-probe', 56510)]
SHORT = {'csf': 'CSF, loss gốc', 'lsa': 'LSA-CSF, loss gốc'}
DS = {'isic17': ('ISIC2017', 2150, 1500, 216, 434), 'isic18': ('ISIC2018', 2694, 1886, 270, 538)}
thin = Side(style='thin'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H1 = PatternFill('solid', fgColor='DDEBF7'); H2 = PatternFill('solid', fgColor='F2F2F2')

def load(path):
    rows = {r['filename']: r for r in csv.DictReader(open(path))}
    return rows

def arr(rows, names, key):
    return np.array([float(rows[n][key]) for n in names])

def best_epoch(run):
    b = [f for f in os.listdir(f'{run}/checkpoints') if f.startswith('best-epoch')]
    return int(b[0].split('epoch')[1].split('-')[0])

def header(ws, r, cols, fill):
    for j, h in enumerate(cols, 1):
        c = ws.cell(r, j, h)
        c.font = Font(bold=True); c.fill = fill; c.border = BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 30

def cell(ws, r, j, v, fmt=None, left=False):
    c = ws.cell(r, j, v)
    c.border = BORDER
    c.alignment = Alignment(horizontal='left' if left else 'center', vertical='center', wrap_text=True)
    if fmt: c.number_format = fmt
    return c

def ttest_row(ws, r, label, v, ref_, names_mask, keys=('dsc', 'iou')):
    n = int(names_mask.sum())
    cell(ws, r, 1, label, left=True); cell(ws, r, 2, n)
    j = 3
    for k in keys:
        a, b = v[k][names_mask], ref_[k][names_mask]
        t = stats.ttest_rel(a, b)
        cell(ws, r, j, 100 * float((a - b).mean()), FMT_D); cell(ws, r, j + 1, float(t.statistic), FMT_T); cell(ws, r, j + 2, float(t.pvalue), FMT_P)
        j += 3

wb = openpyxl.Workbook(); wb.remove(wb.active)
for ds, (nm, n_all, n_tr, n_va, n_te) in DS.items():
    ws = wb.create_sheet(ds)
    for col, w in zip('ABCDEFGHI', (30, 34, 10, 15, 15, 15, 15, 15, 15)):
        ws.column_dimensions[col].width = w
    runs = {k: f'{R}/EGE-UNet-results-exp15/egeunet_{ds}_{suf}_s42' for k, _, suf, _ in ARMS}
    rows = {k: load(f'{runs[k]}/{CSV}') for k in runs}
    names = sorted(rows['base'])
    assert all(sorted(rows[k]) == names for k in rows) and len(names) == n_te
    m = {k: {'dsc': arr(rows[k], names, 'dsc'), 'iou': arr(rows[k], names, 'iou'),
             'tp': arr(rows[k], names, 'tp'), 'fp': arr(rows[k], names, 'fp'), 'fn': arr(rows[k], names, 'fn')} for k in rows}
    for k in m:
        tp, fp, fn = m[k]['tp'], m[k]['fp'], m[k]['fn']
        m[k]['precision'] = np.where(tp + fp > 0, tp / np.maximum(tp + fp, 1), np.nan)
        m[k]['recall'] = np.where(tp + fn > 0, tp / np.maximum(tp + fn, 1), np.nan)
        rr = rows[k]
        m[k]['gt_int'] = np.array([float(rr[n]['gt_interior']) for n in names]); m[k]['fn_int'] = np.array([float(rr[n]['fn_interior']) for n in names])
    area = arr(rows['base'], names, 'area_frac')
    edges = np.quantile(area, [1 / 3, 2 / 3]); bins = np.digitize(area, edges)
    groups = [('Toàn bộ tập test', np.ones(len(names), bool)), ('Tổn thương NHỎ', bins == 0), ('Tổn thương TO', bins == 2)]
    ep = {k: best_epoch(runs[k]) for k in runs}
    old = load(f'{R}/EGE-UNet-results-exp14/egeunet_{ds}_learnable_s42/{CSV}')
    old_dsc = arr(old, names, 'dsc'); t_old = stats.ttest_rel(m['lsa']['dsc'], old_dsc)
    ws['A1'] = f'EXP-15 — {nm} — Model gốc + LSA-CSF (router chọn lọc) so với model gốc và CSF; seed 42, 1 lần train mỗi cấu hình'
    ws['A1'].font = Font(bold=True, size=13)
    ws['A3'] = 'Ghi chú'; ws['A3'].font = Font(bold=True)
    notes = [
        (f'Dữ liệu: {nm} có {n_all} ảnh; chia lại train / val / test = 70 / 10 / 20: {n_tr} ảnh train (giữ nguyên của tác giả) / {n_va} ảnh val / {n_te} ảnh test. '
         f'Val và test lấy từ bản val của tác giả, chia ngẫu nhiên phân tầng theo kích thước tổn thương (seed 42). Ảnh 256×256.', 34),
        (f'Chia nhóm theo diện tích tổn thương thật (tỉ lệ pixel của ảnh), 3 nhóm bằng nhau (tertile): NHỎ < {100 * edges[0]:.2f}% ({int((bins == 0).sum())} ảnh) · '
         f'vừa ({int((bins == 1).sum())} ảnh) · TO ≥ {100 * edges[1]:.2f}% ({int((bins == 2).sum())} ảnh).', 19),
        ('CSF (cross-stage fusion): feature map của cả 5 tầng encoder được chiếu về 8 kênh (conv 1×1 + GroupNorm + GELU), đưa về cùng kích thước, cộng có trọng số học được, '
         'qua conv 1×1 (khởi tạo 0) rồi cộng vào feature của 3 tầng decoder dec3, dec4, dec5 (sau GAB). Thêm 1.591 tham số (53.374 → 54.965).', 34),
        ('LSA-CSF: giữ phần chiếu 8 kênh của CSF; chia thành 2 nhánh (expert): Detail = tầng encoder 1–3 ở lưới 128, Completion = tầng encoder 3–5 ở lưới 32 '
         '(mỗi nhánh: depthwise conv 3×3 → conv 1×1 về 8 kênh → GroupNorm → GELU). Ở mỗi tầng dec3, dec4, dec5 (sau GAB), router (conv 1×1 + softmax) chọn tại từng pixel '
         'tỉ trọng Detail / Completion / không dùng, từ feature decoder + xác suất lesion của đầu ra deep supervision tầng đó (giá trị tại chỗ, trung bình 2 cỡ cửa sổ, trung bình toàn ảnh). '
         'Router không có nhãn riêng. Thêm 3.136 tham số so với model gốc (53.374 → 56.510).', 64),
        (f'Mọi run: seed 42, 300 epoch, loss gốc (BCE + Dice, có deep supervision), chọn checkpoint theo val loss trên tập val ({n_va} ảnh), đánh giá trên tập test riêng. '
         'CSF và LSA-CSF đều so với cùng một model gốc train cùng đợt. DSC, mIoU = trung bình theo ảnh ± độ lệch chuẩn; pooled = tính trên tổng pixel của cả tập.', 34),
        (f'Run: model gốc = egeunet_{ds}_learnable_s42 (best epoch {ep["base"]}); CSF = egeunet_{ds}_learnable_fuse-sum-shallow3-d8_s42 (best epoch {ep["csf"]}); '
         f'LSA-CSF = egeunet_{ds}_learnable_fuse-sum-shallow3-d8_lsa-route-probe_s42 (best epoch {ep["lsa"]}).', 34),
        (f'Tham khảo: lần train trước của model gốc (EXP-14, cùng cấu hình, cùng seed) đạt DSC {100 * old_dsc.mean():.2f} trên cùng tập test; '
         f'LSA-CSF so với lần đó: Δ DSC {100 * (m["lsa"]["dsc"] - old_dsc).mean():+.2f} (p = {t_old.pvalue:.4f}). Hai lần train cùng một model có thể chênh nhau đến vài phần mười điểm DSC.', 34),
    ]
    for i, (txt, h) in enumerate(notes):
        r = 4 + i
        ws.cell(r, 1, txt).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        ws.row_dimensions[r].height = h
    r = 4 + len(notes) + 1
    for gi, (gname, mask) in enumerate(groups):
        n = int(mask.sum())
        ws.cell(r, 1, f'Bảng {gi + 1} — {gname} (n = {n} ảnh)').font = Font(bold=True, size=12); r += 1
        cols = ['Model', 'Loss', 'Params', 'DSC (%)', 'mIoU (%)'] + (['mIoU pooled (%)', 'DSC pooled (%)'] if gi == 0 else [])
        header(ws, r, cols, H1); r += 1
        for k, label, _, params in ARMS:
            d, u = 100 * m[k]['dsc'][mask], 100 * m[k]['iou'][mask]
            cell(ws, r, 1, label, left=True); cell(ws, r, 2, 'loss gốc (BCE + Dice)', left=True); cell(ws, r, 3, params)
            cell(ws, r, 4, f'{d.mean():.2f} ± {d.std(ddof=1):.2f}'); cell(ws, r, 5, f'{u.mean():.2f} ± {u.std(ddof=1):.2f}')
            if gi == 0:
                tp, fp, fn = m[k]['tp'].sum(), m[k]['fp'].sum(), m[k]['fn'].sum()
                cell(ws, r, 6, round(100 * tp / (tp + fp + fn), 2), FMT_POOL); cell(ws, r, 7, round(100 * 2 * tp / (2 * tp + fp + fn), 2), FMT_POOL)
            r += 1
        r += 1
        ws.cell(r, 1, 'Paired t-test theo ảnh so với model gốc (Δ = biến thể − model gốc)').font = Font(bold=True); r += 1
        header(ws, r, ['Biến thể', 'n', 'Δ DSC (%)', 't', 'p', 'Δ mIoU (%)', 't', 'p'], H2); r += 1
        for k in ('csf', 'lsa'):
            ttest_row(ws, r, SHORT[k], m[k], m['base'], mask); r += 1
        r += 1
        ws.cell(r, 1, 'Paired t-test theo ảnh LSA-CSF so với CSF (Δ = LSA-CSF − CSF)').font = Font(bold=True); r += 1
        header(ws, r, ['Biến thể', 'n', 'Δ DSC (%)', 't', 'p', 'Δ mIoU (%)', 't', 'p'], H2); r += 1
        ttest_row(ws, r, 'LSA-CSF so với CSF', m['lsa'], m['csf'], mask); r += 2
    # Bảng 4: precision / recall / interior miss on the whole test set
    ws.cell(r, 1, f'Bảng 4 — Precision, recall và phần bên trong tổn thương bị bỏ sót, toàn bộ tập test (n = {len(names)} ảnh)').font = Font(bold=True, size=12); r += 1
    header(ws, r, ['Model', 'Loss', 'Params', 'Precision (%)', 'Recall (%)', 'Bỏ sót phần bên trong (%)'], H1); r += 1
    for k, label, _, params in ARMS:
        cell(ws, r, 1, label, left=True); cell(ws, r, 2, 'loss gốc (BCE + Dice)', left=True); cell(ws, r, 3, params)
        cell(ws, r, 4, round(100 * float(np.nanmean(m[k]['precision'])), 2), FMT_POOL); cell(ws, r, 5, round(100 * float(np.nanmean(m[k]['recall'])), 2), FMT_POOL)
        cell(ws, r, 6, round(100 * m[k]['fn_int'].sum() / m[k]['gt_int'].sum(), 2), FMT_POOL); r += 1
    r += 1
    ws.cell(r, 1, 'Paired t-test theo ảnh so với model gốc (Δ = biến thể − model gốc)').font = Font(bold=True); r += 1
    header(ws, r, ['Biến thể', 'n', 'Δ Precision (%)', 't', 'p', 'Δ Recall (%)', 't', 'p'], H2); r += 1
    for k in ('csf', 'lsa'):
        ok = np.isfinite(m[k]['precision']) & np.isfinite(m['base']['precision']) & np.isfinite(m[k]['recall']) & np.isfinite(m['base']['recall'])
        ttest_row(ws, r, SHORT[k], m[k], m['base'], ok, keys=('precision', 'recall')); r += 1
    r += 1
    note = ws.cell(r, 1, 'Bỏ sót phần bên trong = tỉ lệ pixel tổn thương có độ sâu chuẩn hóa ≥ 0,15 (khoảng cách tới nền / khoảng cách lớn nhất trong tổn thương của ảnh đó) mà model không tô, cộng trên cả tập test.')
    note.alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9); ws.row_dimensions[r].height = 34
wb.save(OUT)
print('wrote', OUT)
